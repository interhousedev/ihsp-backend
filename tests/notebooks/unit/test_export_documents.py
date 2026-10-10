"""Office exporters: each file is opened again with its own library."""

import csv
import io
from typing import Any

from app.modules.notebooks.application.export.docx import markdown_to_docx
from app.modules.notebooks.application.export.pptx import deck_to_pptx
from app.modules.notebooks.application.export.render import (
    DocumentExporter,
    default_format,
    export_artifact,
    export_as,
    file_name,
    parse_slides,
    parse_table,
)
from app.modules.notebooks.application.export.xlsx import data_table_to_xlsx
from app.modules.notebooks.domain.enums import ArtifactKind

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

MARKDOWN = """# Cells

Intro with **bold** and $E=mc^2$.

## Parts

- nucleus
- membrane

1. first
2. second

| Part | Job |
|------|-----|
| ribosome | builds |
"""


def test_docx_keeps_headings_lists_tables_and_latex() -> None:
    """The report becomes headings, list items, a table and plain LaTeX text."""
    document = Document(io.BytesIO(markdown_to_docx("Biology", MARKDOWN)))

    styles = [(p.style.name, p.text) for p in document.paragraphs]
    assert ("Title", "Biology") in styles
    assert ("Heading 1", "Cells") in styles
    assert ("Heading 2", "Parts") in styles
    assert ("List Bullet", "nucleus") in styles
    assert ("List Number", "second") in styles
    assert any("$E=mc^2$" in text for _, text in styles)
    table = document.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Part", "Job"]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is True
    assert [c.text for c in table.rows[1].cells] == ["ribosome", "builds"]


def test_xlsx_has_bold_header_cells_and_widths() -> None:
    """A data table gets its labels as a bold header and fitted columns."""
    content: dict[str, Any] = {
        "title": "Organelles",
        "columns": [{"key": "name", "label": "Name"}, {"key": "job", "label": "Job"}],
        "rows": [
            {"cells": {"name": "Ribosome", "job": "Builds proteins in the cell"}},
            {"cells": {"name": "Nucleus", "job": "Stores DNA"}},
        ],
    }

    sheet = load_workbook(io.BytesIO(data_table_to_xlsx(content))).active

    assert sheet.title == "Organelles"
    assert [c.value for c in sheet[1]] == ["Name", "Job"]
    assert all(c.font.bold for c in sheet[1])
    assert [c.value for c in sheet[3]] == ["Nucleus", "Stores DNA"]
    assert sheet.max_row == 3
    assert sheet.column_dimensions["B"].width > sheet.column_dimensions["A"].width


def test_pptx_has_a_slide_per_item_with_bullets_and_notes() -> None:
    """A title slide is added; bullets and speaker notes survive."""
    deck: dict[str, Any] = {
        "title": "Cells",
        "subtitle": "Grade 9",
        "slides": [
            {
                "layout": "bullets",
                "title": "Parts",
                "bullets": ["nucleus", "membrane"],
                "notes": "Say it slowly.",
            },
            {
                "layout": "table",
                "title": "Jobs",
                "table": {"columns": ["Part", "Job"], "rows": [["ribosome", "builds"]]},
            },
        ],
    }

    presentation = Presentation(io.BytesIO(deck_to_pptx(deck)))

    assert len(presentation.slides) == 3
    first, second, third = presentation.slides
    assert first.shapes.title.text == "Cells"
    assert second.placeholders[1].text_frame.text == "nucleus\nmembrane"
    assert second.notes_slide.notes_text_frame.text == "Say it slowly."
    assert third.shapes.title.text == "Jobs"
    assert third.shapes[1].table.cell(1, 1).text == "builds"


def test_export_artifact_picks_the_format_of_the_kind() -> None:
    """A report is a .docx whose name comes from the title."""
    doc = export_artifact(
        ArtifactKind.REPORT, "Cells: a / b", {"title": "Cells", "markdown": "# Hi\n"}
    )

    assert doc.name == "Cells a b.docx"
    assert Document(io.BytesIO(doc.data)).paragraphs[1].text == "Hi"


def test_assistant_text_becomes_a_table_and_slides() -> None:
    """CSV and Markdown tables, ``---`` and heading separated slides."""
    assert parse_table("a;b\n1;2") == (["a", "b"], [["1", "2"]])
    assert parse_table("| a | b |\n|---|---|\n| 1 | 2 |") == (["a", "b"], [["1", "2"]])
    slides = parse_slides("# One\n- x\nNotes: hi\n---\n# Two\n- y\n# Three\n- z")
    assert [s["title"] for s in slides] == ["One", "Two", "Three"]
    assert slides[0]["notes"] == "hi"

    exporter = DocumentExporter()
    deck = Presentation(
        io.BytesIO(exporter.render("pptx", "Deck", "# A\n- x\n---\n# B\n- y").data)
    )
    assert len(deck.slides) == 3
    sheet = load_workbook(
        io.BytesIO(exporter.render("xlsx", "T", "a,b\n1,2").data)
    ).active
    assert [c.value for c in sheet[1]] == ["a", "b"]


def test_file_name_drops_unsafe_characters() -> None:
    """Slashes and colons never reach the storage key."""
    assert file_name('a/b:c"', "xlsx") == "a b c.xlsx"
    assert file_name("", "pptx") == "document.pptx"


QUIZ: dict[str, Any] = {
    "questions": [
        {
            "prompt": "Which organelle makes energy?",
            "options": [
                {"id": "a", "text": "Nucleus"},
                {"id": "b", "text": "Mitochondrion"},
            ],
            "correct": ["b"],
            "explanation": "It runs respiration.",
        },
        {
            "prompt": "Pick the membranes",
            "options": [
                {"id": "a", "text": "Plasma"},
                {"id": "b", "text": "Rock"},
                {"id": "c", "text": "Nuclear"},
            ],
            "correct": ["a", "c"],
        },
    ]
}


def test_quiz_docx_lists_questions_then_an_answer_key() -> None:
    """The quiz is a .docx: questions with lettered options, the key at the end."""
    rendered = export_as(ArtifactKind.QUIZ, "Cells quiz", QUIZ, "docx")

    texts = [p.text for p in Document(io.BytesIO(rendered.data)).paragraphs]
    assert rendered.name == "Cells quiz.docx"
    assert "1. Which organelle makes energy?" in texts
    assert "B) Mitochondrion" in texts
    key_at = texts.index("Answer key")
    assert key_at > texts.index("C) Nuclear")
    assert texts[key_at + 1 :] == [
        "1. B",
        "It runs respiration.",
        "2. A, C",
    ]


def test_flashcards_csv_is_front_semicolon_back() -> None:
    """Flash cards are a header-less ``front;back`` CSV Anki can import."""
    content = {
        "cards": [
            {"id": "1", "front": "hola", "back": "hello"},
            {"id": "2", "front": "a; b", "back": "line\nbreak"},
        ]
    }

    rendered = export_as(ArtifactKind.FLASHCARDS, "Verbs", content, "csv")

    assert rendered.name == "Verbs.csv"
    assert rendered.content_type.startswith("text/csv")
    rows = list(csv.reader(io.StringIO(rendered.data.decode()), delimiter=";"))
    assert rows == [["hola", "hello"], ["a; b", "line\nbreak"]]


def test_default_formats_per_kind() -> None:
    """One native file per kind; pictures have only the Markdown fallback."""
    assert default_format(ArtifactKind.REPORT) == "docx"
    assert default_format(ArtifactKind.SLIDES) == "pptx"
    assert default_format(ArtifactKind.DATA_TABLE) == "xlsx"
    assert default_format(ArtifactKind.QUIZ) == "docx"
    assert default_format(ArtifactKind.FLASHCARDS) == "csv"
    assert default_format(ArtifactKind.MIND_MAP) == "md"

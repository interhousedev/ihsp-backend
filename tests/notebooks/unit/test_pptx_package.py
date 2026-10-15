"""The Studio .pptx is a valid OOXML package and downloads byte for byte."""

import io
import posixpath
import zipfile
from types import SimpleNamespace
from typing import Any
from uuid import uuid4

from app.interfaces.http.dependencies import get_actor
from app.interfaces.http.v1.notebooks.deps import get_notebooks
from app.interfaces.http.v1.notebooks.handlers.download_artifact import (
    download_artifact,
)
from app.modules.notebooks.application.export.pptx import build_pptx
from app.modules.notebooks.domain.enums import ArtifactKind, ArtifactStatus

import pytest
from fastapi import FastAPI, Response
from httpx import ASGITransport, AsyncClient
from lxml import etree
from pptx import Presentation

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"

DECK: dict[str, Any] = {
    "title": 'Клетка & «ядро» "<b>" 🧬',
    "subtitle": "Биология\x01 — 9 класс",
    "slides": [
        {"layout": "title", "title": "Клетка & ядро", "bullets": [], "notes": ""},
        {"layout": "agenda", "title": "План", "bullets": ["Ядро", "Мембрана", "Итог"]},
        {
            "layout": "bullets",
            "title": "Мембрана <внешняя>",
            "bullets": ["A & B", "эмодзи 🎓", "bad\x00\x1b￾\ud83d char"],
            "body": 'Цитата "в кавычках"',
            "notes": "Скажите: «привет»\x0c",
        },
        {
            "layout": "two_columns",
            "title": "Сравнение",
            "left": {"heading": "Растения", "bullets": ["Хлоропласты"]},
            "right": {"heading": "Животные", "bullets": ["Центриоли"]},
        },
        {
            "layout": "quote",
            "title": "",
            "quote": {"text": "Omnis cellula e cellula", "attribution": "Вирхов"},
        },
        {
            "layout": "table",
            "title": "Органоиды",
            "table": {"columns": ["Часть", "Роль"], "rows": [["Рибосома", "белок"]]},
        },
        {"layout": "section", "title": "Итоги", "bullets": []},
        {"layout": "summary", "title": "Итог", "bullets": ["Клетка — основа"]},
    ],
}


def assert_valid_package(data: bytes) -> None:
    """python-pptx opens it, every part parses, types and relationships resolve."""
    Presentation(io.BytesIO(data))
    package = zipfile.ZipFile(io.BytesIO(data))
    assert package.testzip() is None
    names = package.namelist()
    assert len(names) == len(set(names))
    for name in names:
        if name.endswith((".xml", ".rels")):
            etree.fromstring(package.read(name))
    types = etree.fromstring(package.read("[Content_Types].xml"))
    defaults = {e.get("Extension").lower() for e in types.iter(f"{_CT}Default")}
    overrides = {e.get("PartName") for e in types.iter(f"{_CT}Override")}
    for name in names:
        if name != "[Content_Types].xml":
            ext = name.rsplit(".", 1)[-1].lower()
            assert f"/{name}" in overrides or ext in defaults, name
    assert all(part[1:] in names for part in overrides)
    for name in (n for n in names if n.endswith(".rels")):
        base = posixpath.dirname(posixpath.dirname(name))
        for rel in etree.fromstring(package.read(name)):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            path = (
                target[1:]
                if target.startswith("/")
                else posixpath.normpath(posixpath.join(base, target))
            )
            assert path in names, (name, target)


def test_deck_is_a_valid_package_with_every_slide() -> None:
    data = build_pptx(DECK)
    assert_valid_package(data)
    slides = Presentation(io.BytesIO(data)).slides
    assert len(slides) == len(DECK["slides"])
    texts = " ".join(
        shape.text_frame.text
        for slide in slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "A & B" in texts and "🎓" in texts and "bad char" in texts


def test_a_deck_without_a_title_slide_gets_one() -> None:
    data = build_pptx({"title": "T", "subtitle": "x\ud800", "slides": []})
    assert_valid_package(data)
    assert len(Presentation(io.BytesIO(data)).slides) == 1


@pytest.mark.asyncio
async def test_download_returns_the_exact_bytes_as_pptx() -> None:
    artifact = SimpleNamespace(
        kind=ArtifactKind.SLIDES,
        status=ArtifactStatus.READY,
        title="Клетка",
        content=DECK,
    )

    async def execute(*_: Any) -> Any:
        return SimpleNamespace(artifact=artifact)

    services = SimpleNamespace(get_artifact=SimpleNamespace(execute=execute))
    app = FastAPI()
    app.add_api_route(
        "/n/{notebook_id}/a/{artifact_id}",
        download_artifact,
        methods=["GET"],
        response_class=Response,
    )
    app.dependency_overrides[get_notebooks] = lambda: services
    app.dependency_overrides[get_actor] = lambda: SimpleNamespace()
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://t"
    ) as client:
        response = await client.get(f"/n/{uuid4()}/a/{uuid4()}")
    assert response.status_code == 200
    assert response.headers["content-type"] == PPTX_TYPE
    disposition = response.headers["content-disposition"]
    assert 'filename="file.pptx"' in disposition
    assert "filename*=UTF-8''%D0%9A%D0%BB%D0%B5%D1%82%D0%BA%D0%B0.pptx" in disposition
    assert int(response.headers["content-length"]) == len(response.content)
    assert response.content[:2] == b"PK"
    assert_valid_package(response.content)

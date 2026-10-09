"""Realistic sources for the slide deck tests and the eval harness.

Five sample files of the kinds students add to a notebook: a long English PDF, an
English lesson note, a Russian DOCX with headings and a list, Russian teacher's notes
saved in Windows-1251 and an Uzbek (Latin) text. The texts live in `sources/`; this
module turns them into real file bytes (PDF, DOCX, encoded TXT), runs the real
extractors and the chunker, and gives stored-like `Source` and `Passage` entities.
`recorded/` holds model answers recorded from real runs of
`scripts/nb_eval_slides.py` (used with the scripted fake model).

Used by `tests/notebooks/unit/domain/test_slides_realistic.py`,
`tests/notebooks/unit/studio/test_slides_realistic.py` and the eval harness.
"""

import io
import json
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal
from uuid import UUID

from app.modules.files.infrastructure.extraction.docx_text import extract_docx
from app.modules.files.infrastructure.extraction.pdf_text import extract_pdf
from app.modules.files.infrastructure.extraction.plain_text import extract_plain
from app.modules.files.ports.text_extractor import ExtractedText
from app.modules.notebooks.domain.chunker import ChunkResult, chunk_source
from app.modules.notebooks.domain.entities import Passage, Source
from app.modules.notebooks.domain.enums import SourceKind, SourceStatus
from app.shared.ids import new_id
from tests.files.helpers import make_pdf

from docx import Document

HERE = Path(__file__).parent
SOURCES_DIR = HERE / "sources"
RECORDED_DIR = HERE / "recorded"
PAGE_BREAK = "\n--- page ---\n"
NOW = datetime(2026, 10, 1, tzinfo=UTC)

FileKind = Literal["pdf", "docx", "txt"]


@dataclass(frozen=True, slots=True)
class SampleSource:
    """One sample file: its title in the notebook, its kind and where its text is."""

    key: str
    title: str
    kind: FileKind
    file: str
    encoding: str = "utf-8"


SAMPLES: dict[str, SampleSource] = {
    sample.key: sample
    for sample in (
        SampleSource(
            "industrial_revolution",
            "The Industrial Revolution.pdf",
            "pdf",
            "industrial_revolution.pdf.txt",
        ),
        SampleSource(
            "factory_acts", "Factory Acts (lesson notes).txt", "txt", "factory_acts.txt"
        ),
        SampleSource(
            "peter_reforms", "Реформы Петра I.docx", "docx", "peter_reforms.docx.md"
        ),
        SampleSource(
            "peter_dates",
            "Пётр I — даты.txt",
            "txt",
            "peter_dates.cp1251.txt",
            encoding="cp1251",
        ),
        SampleSource("amir_temur", "Amir Temur.txt", "txt", "amir_temur.txt"),
    )
}


def source_text(key: str) -> str:
    """The text a sample file is made from (pages split by `PAGE_BREAK` for PDFs)."""
    return (SOURCES_DIR / SAMPLES[key].file).read_text(encoding="utf-8")


def _docx_bytes(markup: str) -> bytes:
    """A DOCX built from light markup: ``#`` title, ``##`` heading, ``-`` list item."""
    document = Document()
    for block in markup.strip().split("\n\n"):
        for line in block.split("\n"):
            if line.startswith("## "):
                document.add_heading(line[3:], 1)
            elif line.startswith("# "):
                document.add_heading(line[2:], 0)
            elif line.startswith("- "):
                document.add_paragraph(line[2:], style="List Bullet")
            else:
                document.add_paragraph(line)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def file_bytes(key: str) -> bytes:
    """The sample as the student uploads it."""
    sample = SAMPLES[key]
    text = source_text(key)
    if sample.kind == "pdf":
        return make_pdf(*(page.strip() for page in text.split(PAGE_BREAK)))
    if sample.kind == "docx":
        return _docx_bytes(text)
    return text.encode(sample.encoding)


def extracted(key: str) -> ExtractedText:
    """What the files module stores for the sample (the real extractors)."""
    kind = SAMPLES[key].kind
    data = file_bytes(key)
    if kind == "pdf":
        return extract_pdf(data)
    if kind == "docx":
        return extract_docx(data)
    return extract_plain(data)


def chunked(key: str) -> ChunkResult:
    """The sample after normalisation and chunking, as the source pipeline does it."""
    text = extracted(key)
    return chunk_source(text.text, page_starts=text.page_starts)


def stored(
    keys: Sequence[str], notebook_id: UUID | None = None
) -> tuple[list[Source], dict[UUID, list[Passage]]]:
    """Ready `Source` entities (in the order of `keys`) and their passages."""
    notebook = notebook_id or new_id()
    sources: list[Source] = []
    passages: dict[UUID, list[Passage]] = {}
    for position, key in enumerate(keys):
        result = chunked(key)
        source = Source(
            id=new_id(),
            notebook_id=notebook,
            owner_id=new_id(),
            kind=SourceKind.FILE,
            title=SAMPLES[key].title,
            position=position,
            created_at=NOW,
            updated_at=NOW,
            status=SourceStatus.READY,
            char_count=result.char_count,
            token_estimate=result.token_estimate,
            passage_count=len(result.passages),
        )
        sources.append(source)
        passages[source.id] = [
            Passage(
                id=new_id(),
                source_id=source.id,
                notebook_id=notebook,
                seq=draft.seq,
                text=draft.text,
                char_start=draft.char_start,
                char_end=draft.char_end,
                token_estimate=draft.token_estimate,
                page=draft.page,
                heading=draft.heading,
            )
            for draft in result.passages
        ]
    return sources, passages


@dataclass
class MemoryPassages:
    """`list_for_sources` and a word-overlap `search` over stored passages."""

    by_source: dict[UUID, list[Passage]]

    async def list_for_sources(self, ids: Sequence[UUID]) -> list[Passage]:
        """The passages of the sources, in source order then seq."""
        return [p for sid in ids for p in self.by_source.get(sid, [])]

    async def search(
        self, notebook_id: UUID, ids: Sequence[UUID], query: str, limit: int
    ) -> list[Passage]:
        """Passages that share the most words (4+ letters) with the query."""
        words = {w.casefold() for w in re.findall(r"\w{4,}", query)}
        scored = [
            (sum(w in p.text.casefold() for w in words), p)
            for sid in ids
            for p in self.by_source.get(sid, [])
        ]
        ranked = sorted((s for s in scored if s[0]), key=lambda s: -s[0])
        return [p for _, p in ranked][:limit]


@dataclass
class MemoryUow:
    """The part of the notebooks unit of work that the Studio context needs."""

    passages: MemoryPassages


def recorded(name: str) -> dict[str, Any]:
    """A model answer recorded by the eval harness (`recorded/<name>.json`)."""
    data = json.loads((RECORDED_DIR / f"{name}.json").read_text(encoding="utf-8"))
    assert isinstance(data, dict)
    return data

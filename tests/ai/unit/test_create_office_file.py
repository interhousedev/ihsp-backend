"""create_presentation / create_document / create_spreadsheet: build, store, card."""

import io
import json
from typing import Any

from app.modules.files.domain.errors import QuotaExceededError
from app.modules.notebooks.application.export.render import DocumentExporter
from app.modules.notebooks.application.export.structured import (
    blocks_to_docx,
    presentation_to_pptx,
    sheets_to_xlsx,
)
from app.shared.contracts.files import FilePurpose, FileScope
from tests.ai.conftest import AiHarness
from tests.fakes.ai.llm import text_reply, tool_reply

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

SLIDES = [{"title": "Клетка", "bullets": ["Ядро\x07", "Мембрана"], "notes": "Скажи"}]
BLOCKS = [
    {"type": "heading", "text": "Введение", "level": 1},
    {"type": "paragraph", "text": "Текст\x00 абзаца"},
    {"type": "bullets", "items": ["один", "два"]},
    {"type": "table", "columns": ["Имя", "Балл"], "rows": [["Аня", 5]]},
]
SHEETS = [
    {"name": "Оценки", "columns": ["Имя", "Балл"], "rows": [["Аня", 5], ["=1+1", 4.5]]}
]


def test_presentation_builder_reopens_with_cyrillic_and_notes() -> None:
    """A title slide plus one slide per part; control characters are gone."""
    deck = Presentation(io.BytesIO(presentation_to_pptx("Биология", SLIDES)))

    slides = list(deck.slides)
    assert len(slides) == 2 and slides[0].shapes.title.text == "Биология"
    body = slides[1].placeholders[1].text_frame.text
    assert slides[1].shapes.title.text == "Клетка" and body == "Ядро\nМембрана"
    assert slides[1].notes_slide.notes_text_frame.text == "Скажи"


def test_document_builder_reopens_with_every_block() -> None:
    """Heading, paragraph, bullets and a table survive the round trip."""
    doc = Document(io.BytesIO(blocks_to_docx("Отчёт", BLOCKS)))

    texts = [p.text for p in doc.paragraphs]
    assert texts[:5] == ["Отчёт", "Введение", "Текст абзаца", "один", "два"]
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Имя", "Балл"]
    assert [c.text for c in table.rows[1].cells] == ["Аня", "5"]


def test_spreadsheet_builder_keeps_numbers_and_defuses_formulas() -> None:
    """Numbers stay numbers; a text starting with '=' is not a formula."""
    book = load_workbook(io.BytesIO(sheets_to_xlsx(SHEETS)))

    sheet = book["Оценки"]
    assert [c.value for c in sheet[1]] == ["Имя", "Балл"]
    assert [c.value for c in sheet[2]] == ["Аня", 5]
    assert sheet["A3"].value == "'=1+1" and sheet["B3"].value == 4.5


@pytest.mark.parametrize(
    ("tool", "args", "kind"),
    [
        ("create_presentation", {"title": "Биология", "slides": SLIDES}, "pptx"),
        ("create_document", {"title": "Отчёт", "blocks": BLOCKS}, "docx"),
        ("create_spreadsheet", {"title": "Оценки", "sheets": SHEETS}, "xlsx"),
    ],
)
async def test_tool_stores_a_file_for_the_actor_and_returns_its_id(
    h: AiHarness, tool: str, args: dict[str, Any], kind: str
) -> None:
    """The file is the student's own, the model gets its id, the chat a card."""
    h.services.bindings.documents = DocumentExporter()
    chat = await h.new_chat()
    h.llm.push(tool_reply((tool, args)), text_reply("Готово."))

    result = await h.say(chat.id, "сделай файл")

    call = result.new_messages[1].tool_calls[0]
    assert call.status == "done" and call.document is not None
    assert call.document.kind == kind and call.document.name.endswith("." + kind)
    ref = h.files.files[call.document.file_id]
    assert ref.owner_id == h.student.user_id and ref.scope is FileScope.STUDENT
    assert ref.purpose is FilePurpose.AI_UPLOAD
    assert h.files.stored[ref.id][:2] == b"PK"
    sent = json.loads(h.llm.requests[-1]["messages"][-1]["content"])
    assert sent["untrusted_data"]["file_id"] == str(ref.id)


async def test_too_many_slides_are_rejected(h: AiHarness) -> None:
    """More than 30 slides: an argument error, nothing stored."""
    h.services.bindings.documents = DocumentExporter()
    chat = await h.new_chat()
    args = {"title": "T", "slides": [{"title": "s", "bullets": []}] * 31}
    h.llm.push(tool_reply(("create_presentation", args)), text_reply("Sorry."))

    result = await h.say(chat.id, "deck")

    assert result.new_messages[1].tool_calls[0].status == "failed"
    assert not h.files.stored


async def test_quota_exceeded_is_reported_to_the_model(
    h: AiHarness, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A full quota fails the call; the model sees the files quota code."""
    h.services.bindings.documents = DocumentExporter()

    async def full(**_: Any) -> Any:
        raise QuotaExceededError(params={"used_bytes": 1, "quota_bytes": 1, "size": 1})

    monkeypatch.setattr(h.files, "store_bytes", full)
    chat = await h.new_chat()
    args = {"title": "T", "sheets": SHEETS}
    h.llm.push(tool_reply(("create_spreadsheet", args)), text_reply("Full."))

    result = await h.say(chat.id, "table")

    assert result.new_messages[1].tool_calls[0].status == "failed"
    assert "files/quota-exceeded" in h.llm.requests[-1]["messages"][-1]["content"]

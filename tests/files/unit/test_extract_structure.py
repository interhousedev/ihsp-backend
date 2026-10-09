"""FILE-7, NB-2: extraction keeps the structure and every visible word of a file.

DOCX headings become Markdown headings (the notebooks chunker cuts passages on them),
paragraphs are blocks, list items start with "- ", text inside content controls and
tracked insertions is read, merged cells are written once. Plain text files saved in
UTF-16 or Windows-1251 are decoded instead of turning into replacement characters.
"""

import io

from app.modules.files.infrastructure.extraction.docx_text import extract_docx
from app.modules.files.infrastructure.extraction.plain_text import extract_plain
from app.modules.notebooks.domain.chunker import chunk_source

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

W = nsdecls("w")


def saved(document: object) -> bytes:
    """The bytes of a python-docx document."""
    out = io.BytesIO()
    document.save(out)  # type: ignore[attr-defined]
    return out.getvalue()


def lesson_docx() -> bytes:
    """A Russian lesson: title, two sections with paragraphs, a list and a table."""
    document = Document()
    document.add_heading("Фотосинтез", 0)
    document.add_heading("Световая фаза", 1)
    document.add_paragraph("Хлорофилл поглощает свет и запасает энергию в АТФ.")
    document.add_heading("Темновая фаза", 1)
    document.add_paragraph("Углекислый газ", style="List Bullet")
    document.add_paragraph("Глюкоза", style="List Bullet")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Стадия"
    table.cell(0, 2).text = "Где"
    for column, value in enumerate(["Цикл", "Кальвина", "Строма"]):
        table.cell(1, column).text = value
    return saved(document)


def test_docx_headings_paragraphs_and_lists_keep_their_structure() -> None:
    """Title -> #, Heading 1 -> ##, list items -> "- ", blocks split by blank lines."""
    text = extract_docx(lesson_docx()).text

    assert text.split("\n\n") == [
        "# Фотосинтез",
        "## Световая фаза",
        "Хлорофилл поглощает свет и запасает энергию в АТФ.",
        "## Темновая фаза",
        "- Углекислый газ",
        "- Глюкоза",
        "Стадия\tГде\nЦикл\tКальвина\tСтрома",
    ]


def test_docx_sections_become_passages_with_their_heading() -> None:
    """NB-2: the chunker sees the DOCX headings, so passages follow the sections."""
    passages = chunk_source(extract_docx(lesson_docx()).text).passages

    headings = [p.heading for p in passages]
    assert "Световая фаза" in headings
    assert headings[-1] == "Темновая фаза"
    assert passages[-1].text.startswith("## Темновая фаза")


def test_docx_reads_content_controls_and_tracked_insertions() -> None:
    """Text python-docx's Paragraph.text misses is not lost."""
    document = Document()
    body = document.element.body
    block_sdt = parse_xml(
        f"<w:sdt {W}><w:sdtContent><w:p><w:r><w:t>Inside a block control</w:t>"
        "</w:r></w:p></w:sdtContent></w:sdt>"
    )
    paragraph = parse_xml(
        f"<w:p {W}><w:r><w:t xml:space='preserve'>Start </w:t></w:r>"
        "<w:sdt><w:sdtContent><w:r><w:t xml:space='preserve'>inline </w:t></w:r>"
        "</w:sdtContent></w:sdt><w:ins w:id='1' w:author='T'><w:r>"
        "<w:t>inserted</w:t></w:r></w:ins><w:del w:id='2' w:author='T'><w:r>"
        "<w:delText>deleted</w:delText></w:r></w:del></w:p>"
    )
    body.insert(0, paragraph)
    body.insert(0, block_sdt)

    text = extract_docx(saved(document)).text

    assert text.split("\n\n") == ["Inside a block control", "Start inline inserted"]


def test_docx_custom_heading_style_by_outline_level() -> None:
    """A paragraph with an outline level is a heading even with a custom style."""
    document = Document()
    document.add_paragraph("Раздел первый")
    document.element.body[0].insert(
        0, parse_xml(f"<w:pPr {W}><w:outlineLvl w:val='0'/></w:pPr>")
    )
    document.add_paragraph("Matn.")

    assert extract_docx(saved(document)).text == "## Раздел первый\n\nMatn."


def test_plain_text_in_windows_1251_is_decoded() -> None:
    """Russian .txt from old Windows editors."""
    text = "Фотосинтез — процесс образования органических веществ. Ёлка."

    assert extract_plain(text.encode("cp1251")).text == text


def test_plain_text_in_utf16_with_a_bom_is_decoded() -> None:
    """Notepad's "Unicode" encoding."""
    text = "Oʻzbekiston — Markaziy Osiyodagi davlat."

    assert extract_plain(text.encode("utf-16")).text == text


def test_utf8_with_one_broken_byte_stays_utf8() -> None:
    """A stray byte does not turn a UTF-8 text into code page 1251."""
    data = "Привет".encode() + b"\xff" + " мир, oʻquvchi".encode()

    assert extract_plain(data).text == "Привет\N{REPLACEMENT CHARACTER} мир, oʻquvchi"

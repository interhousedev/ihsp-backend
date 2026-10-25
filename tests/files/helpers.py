"""Helpers of the files tests: tiny real PDFs, DOCX files and upload steps."""

import base64
import hashlib
import io
import struct
import zipfile
import zlib
from collections.abc import Mapping, Sequence
from typing import Any
from uuid import UUID

from app.modules.files.application.file_views import FileView
from app.modules.files.application.request_upload import UploadTicket
from app.shared.actor import Actor
from app.shared.contracts.files import FilePurpose
from tests.fakes.files.services import FilesFakes

from docx import Document
from pypdf import PdfReader, PdfWriter

PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
EXE_BYTES = b"MZ" + b"\x90" * 64
PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def sha256_b64(data: bytes) -> str:
    """Base64 of the raw SHA-256 digest, as the browser sends it."""
    return base64.b64encode(hashlib.sha256(data).digest()).decode()


def make_pdf(*pages: str) -> bytes:
    """Build a valid PDF whose pages contain the given texts (Helvetica)."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(pages))]
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    font_id = 3 + 2 * len(pages)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for index, text in enumerate(pages):
        content_id = page_ids[index] + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_id} 0 R >>"
            ).encode()
        )
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 20 150 Td ({escaped}) Tj ET".encode()
        objects.append(
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    """Build a DOCX with paragraphs and an optional table."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


async def upload_and_complete(
    fakes: FilesFakes,
    actor: Actor,
    *,
    name: str,
    data: bytes,
    content_type: str,
    purpose: FilePurpose = FilePurpose.LINK,
) -> FileView:
    """Run request_upload, the browser's PUT and complete_upload."""
    services = fakes.services
    ticket: UploadTicket = await services.request_upload.execute(
        actor,
        name=name,
        size=len(data),
        content_type=content_type,
        sha256_b64=sha256_b64(data),
        purpose=purpose,
    )
    fakes.storage.upload(ticket.upload, data)
    return await services.complete_upload.execute(actor, ticket.view.file.id)


async def upload_ready(
    fakes: FilesFakes,
    actor: Actor,
    *,
    name: str,
    data: bytes,
    content_type: str,
    purpose: FilePurpose = FilePurpose.LINK,
) -> UUID:
    """Upload, complete and process a file; return its id."""
    view = await upload_and_complete(
        fakes, actor, name=name, data=data, content_type=content_type, purpose=purpose
    )
    await fakes.services.process_file.execute(view.file.id)
    return view.file.id


def _png_chunk(kind: bytes, payload: bytes) -> bytes:
    body = kind + payload
    return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body))


def real_png() -> bytes:
    """A valid 1x1 PNG (for the real libmagic)."""
    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\xff\x00\x00")
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", header)
        + _png_chunk(b"IDAT", pixels)
        + _png_chunk(b"IEND", b"")
    )


def real_jpeg() -> bytes:
    """The start of a JFIF JPEG (enough for the real libmagic)."""
    jfif = b"\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    return b"\xff\xd8" + jfif + b"\xff\xd9"


EPUB_TYPE = "application/epub+zip"
XHTML_NS = (
    'xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"'
)
type TocItem = tuple[str, str, list[TocItem]]  # title, href from the nav, children


def _nav_list(items: Sequence[TocItem]) -> str:
    rows = "".join(
        f'<li><a href="{href}">{title}</a>{_nav_list(children) if children else ""}'
        "</li>"
        for title, href, children in items
    )
    return f"<ol>{rows}</ol>"


def _ncx_points(items: Sequence[TocItem], counter: list[int]) -> str:
    points = []
    for title, href, children in items:
        counter[0] += 1
        points.append(
            f'<navPoint id="n{counter[0]}" playOrder="{counter[0]}">'
            f"<navLabel><text>{title}</text></navLabel>"
            f'<content src="{href}"/>{_ncx_points(children, counter)}</navPoint>'
        )
    return "".join(points)


def make_epub(
    chapters: Sequence[str],
    *,
    toc: Sequence[TocItem] = (),
    toc_kind: str = "nav",
    folder: str = "OEBPS",
    extra_files: Mapping[str, bytes] | None = None,
    spine_order: Sequence[int] | None = None,
) -> bytes:
    """Build an EPUB whose chapters are the given XHTML bodies (``ch1.xhtml``, ...).

    ``toc`` is written as an EPUB 3 navigation document (``toc_kind="nav"``), an
    EPUB 2 NCX (``"ncx"``) or both (``"both"``); hrefs are relative to the package
    folder. ``spine_order`` lists chapter numbers (0-based) to reorder the spine.
    """
    prefix = f"{folder}/" if folder else ""
    names = [f"ch{index + 1}.xhtml" for index in range(len(chapters))]
    manifest = [
        f'<item id="c{i}" href="{name}" media-type="application/xhtml+xml"/>'
        for i, name in enumerate(names)
    ]
    order = spine_order if spine_order is not None else range(len(chapters))
    spine = "".join(f'<itemref idref="c{i}"/>' for i in order)
    files: dict[str, bytes] = {}
    with_nav = toc and toc_kind in {"nav", "both"}
    with_ncx = toc and toc_kind in {"ncx", "both"}
    if with_nav:
        manifest.append(
            '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" '
            'properties="nav"/>'
        )
        nav = (
            f"<html {XHTML_NS}><head><title>Contents</title></head><body>"
            f'<nav epub:type="toc">{_nav_list(toc)}</nav></body></html>'
        )
        files[f"{prefix}nav.xhtml"] = nav.encode()
    if with_ncx:
        manifest.append(
            '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>'
        )
        ncx = (
            '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">'
            f"<navMap>{_ncx_points(toc, [0])}</navMap></ncx>"
        )
        files[f"{prefix}toc.ncx"] = ncx.encode()
    toc_attr = ' toc="ncx"' if with_ncx else ""
    opf = (
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
        "<metadata/>"
        f"<manifest>{''.join(manifest)}</manifest>"
        f"<spine{toc_attr}>{spine}</spine></package>"
    )
    container = (
        '<container version="1.0" '
        'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
        f'<rootfile full-path="{prefix}content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    for name, body in zip(names, chapters, strict=True):
        files[f"{prefix}{name}"] = (
            f"<html {XHTML_NS}><head><title>t</title></head><body>{body}</body></html>"
        ).encode()
    files = {
        "META-INF/container.xml": container.encode(),
        f"{prefix}content.opf": opf.encode(),
        **files,
        **(extra_files or {}),
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("mimetype", EPUB_TYPE, compress_type=zipfile.ZIP_STORED)
        for path, content in files.items():
            archive.writestr(path, content, compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


def make_pdf_with_outline(
    pages: Sequence[str], outline: Sequence[tuple[str, int, int]]
) -> bytes:
    """A PDF with bookmarks: ``(title, 0-based page, level)``, levels from 0 up."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(*pages))))
    parents: list[Any] = []
    for title, page, level in outline:
        parent = parents[level - 1] if level else None
        item = writer.add_outline_item(title, page, parent=parent)
        del parents[level:]
        parents.append(item)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# - - - Tiny images (structure only: enough for the PDF assembler) - - -


def _exif_orientation(orientation: int, *, little_endian: bool = False) -> bytes:
    """An APP1 segment whose only tag is the orientation."""
    order = "<" if little_endian else ">"
    tiff = (b"II*\x00" if little_endian else b"MM\x00*") + struct.pack(order + "I", 8)
    tiff += struct.pack(order + "H", 1)
    tiff += struct.pack(order + "HHI", 0x0112, 3, 1)
    tiff += struct.pack(order + "HH", orientation, 0) + struct.pack(order + "I", 0)
    body = b"Exif\x00\x00" + tiff
    return b"\xff\xe1" + struct.pack(">H", len(body) + 2) + body


def make_jpeg(
    width: int,
    height: int,
    *,
    components: int = 3,
    marker: int = 0xC0,
    precision: int = 8,
    orientation: int | None = None,
    little_endian: bool = False,
) -> bytes:
    """A JPEG with real headers and a dummy scan (viewers need a real encoder)."""
    out = b"\xff\xd8" + b"\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    if orientation is not None:
        out += _exif_orientation(orientation, little_endian=little_endian)
    parts = b"".join(bytes([n + 1, 0x11, 0]) for n in range(components))
    frame = struct.pack(">BHHB", precision, height, width, components) + parts
    out += bytes([0xFF, marker]) + struct.pack(">H", len(frame) + 2) + frame
    return out + b"\xff\xda\x00\x08\x01\x01\x00\x00?\x00" + b"\x00" * 16 + b"\xff\xd9"


def make_png(
    width: int,
    height: int,
    *,
    color_type: int = 2,
    bit_depth: int = 8,
    interlace: int = 0,
    chunks: Sequence[tuple[bytes, bytes]] = (),
    idat_parts: int = 1,
) -> bytes:
    """A PNG of zeros; ``chunks`` are extra chunks put before the image data."""
    channels = {0: 1, 2: 3, 3: 1, 4: 2, 6: 4}[color_type]
    row = b"\x00" + bytes(width * channels * bit_depth // 8)
    pixels = zlib.compress(row * height)
    size = max(1, len(pixels) // idat_parts)
    pieces = [pixels[i : i + size] for i in range(0, len(pixels), size)]
    header = struct.pack(
        ">IIBBBBB", width, height, bit_depth, color_type, 0, 0, interlace
    )
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", header)
        + b"".join(_png_chunk(kind, payload) for kind, payload in chunks)
        + b"".join(_png_chunk(b"IDAT", piece) for piece in pieces)
        + _png_chunk(b"IEND", b"")
    )

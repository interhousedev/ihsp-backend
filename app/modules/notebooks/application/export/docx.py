"""Word export: Markdown text to a .docx (headings, lists, tables; LaTeX stays text)."""

import io
import re
from typing import Final

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt
from docx.text.paragraph import Paragraph

_HEADING: Final = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET: Final = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED: Final = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_FENCE: Final = re.compile(r"^\s{0,3}(```|~~~)")
_RULE: Final = re.compile(r"^\s{0,3}([-*_])(\s*\1){2,}\s*$")
_SEPARATOR: Final = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_INLINE: Final = re.compile(r"(\*\*[^*\n]+\*\*|`[^`\n]+`)")


def _runs(paragraph: Paragraph, text: str) -> None:
    """Add the text to the paragraph; ``**bold**`` and `code` get their own runs."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(part)


def _cells(line: str) -> list[str]:
    """Split a Markdown table row into its cells."""
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _table(document: DocxDocument, lines: list[str]) -> None:
    """Write a Markdown table (the separator row is dropped; the header is bold)."""
    rows = [_cells(line) for line in lines if not _SEPARATOR.match(line)]
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            text = row[c] if c < len(row) else ""
            _runs(paragraph, text)
            if r == 0:
                for run in paragraph.runs:
                    run.bold = True


def _code(document: DocxDocument, lines: list[str]) -> None:
    """Write a fenced block as monospace paragraphs."""
    for line in lines:
        run = document.add_paragraph().add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)


def _text_line(document: DocxDocument, line: str, paragraph: list[str]) -> None:
    """One ordinary line: a heading, a list item, a break or part of a paragraph."""
    heading = _HEADING.match(line)
    bullet = _BULLET.match(line)
    numbered = _NUMBERED.match(line)
    if not line.strip() or _RULE.match(line):
        _flush(document, paragraph)
    elif heading:
        _flush(document, paragraph)
        document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
    elif bullet:
        _flush(document, paragraph)
        _runs(document.add_paragraph(style="List Bullet"), bullet.group(2))
    elif numbered:
        _flush(document, paragraph)
        _runs(document.add_paragraph(style="List Number"), numbered.group(2))
    else:
        paragraph.append(line.strip())


def _flush(document: DocxDocument, paragraph: list[str]) -> None:
    """Write the collected lines as one paragraph."""
    if paragraph:
        _runs(document.add_paragraph(), " ".join(paragraph))
        paragraph.clear()


def _add_markdown(document: DocxDocument, markdown: str) -> None:
    """Walk the lines once: fences and tables collect their lines first."""
    lines = markdown.replace("\r\n", "\n").split("\n")
    i = 0
    paragraph: list[str] = []
    while i < len(lines):
        line = lines[i]
        if _FENCE.match(line):
            _flush(document, paragraph)
            end = i + 1
            while end < len(lines) and not _FENCE.match(lines[end]):
                end += 1
            _code(document, lines[i + 1 : end])
            i = end + 1
        elif line.lstrip().startswith("|") and line.count("|") >= 2:
            _flush(document, paragraph)
            end = i
            while end < len(lines) and lines[end].lstrip().startswith("|"):
                end += 1
            _table(document, lines[i:end])
            i = end
        else:
            _text_line(document, line, paragraph)
            i += 1
    _flush(document, paragraph)


def markdown_to_docx(title: str, markdown: str) -> bytes:
    """Return a .docx with the title and the Markdown body."""
    document = Document()
    if title.strip():
        document.add_heading(title.strip(), level=0)
    _add_markdown(document, markdown)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()

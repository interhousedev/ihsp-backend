"""Quiz and flash card exports: a Word test sheet with an answer key, a CSV deck."""

import csv
import io
from typing import Any

from docx import Document
from docx.shared import Pt

CSV_BOM = "﻿"


def _letter(index: int) -> str:
    return chr(ord("A") + index) if index < 26 else str(index + 1)


def _questions(content: dict[str, Any]) -> list[dict[str, Any]]:
    return [q for q in content.get("questions") or [] if isinstance(q, dict)]


def quiz_to_docx(title: str, content: dict[str, Any]) -> bytes:
    """A test sheet: numbered questions with lettered options; the key comes last."""
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(str(content.get("title") or title), level=0)
    key: list[str] = []
    questions = _questions(content)
    for number, question in enumerate(questions, start=1):
        document.add_paragraph().add_run(
            f"{number}. {question.get('prompt', '')}"
        ).bold = True
        correct = {str(c) for c in question.get("correct") or []}
        letters: list[str] = []
        for index, option in enumerate(question.get("options") or []):
            letter = _letter(index)
            document.add_paragraph(f"{letter}) {option.get('text', '')}")
            if str(option.get("id")) in correct:
                letters.append(letter)
        key.append(f"{number}. {', '.join(letters)}")
    document.add_page_break()
    document.add_heading("Answer key", level=1)
    for line, question in zip(key, questions, strict=True):
        document.add_paragraph(line)
        if question.get("explanation"):
            document.add_paragraph(str(question["explanation"]))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def flashcards_to_csv(content: dict[str, Any]) -> bytes:
    """``front;back`` rows without a header: Anki imports the file as it is."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    for card in content.get("cards") or []:
        writer.writerow([str(card.get("front", "")), str(card.get("back", ""))])
    return buffer.getvalue().encode()

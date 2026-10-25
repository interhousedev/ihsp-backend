"""Excel export: a table to an .xlsx (bold header row, fitted column widths)."""

import io
import re
from collections.abc import Sequence
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

MAX_WIDTH: Final = 60
MIN_WIDTH: Final = 8
_BAD_SHEET_CHARS: Final = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(title: str) -> str:
    """Excel sheet names: at most 31 characters, none of the characters ``[]:*?/``."""
    clean = _BAD_SHEET_CHARS.sub(" ", title).strip()[:31].strip()
    return clean or "Sheet1"


def _value(cell: Any) -> Any:
    """Numbers stay numbers; everything else becomes text."""
    if isinstance(cell, bool) or cell is None:
        return "" if cell is None else str(cell)
    if isinstance(cell, int | float):
        return cell
    return str(cell)


def table_to_xlsx(
    title: str, columns: Sequence[str], rows: Sequence[Sequence[Any]]
) -> bytes:
    """Return a workbook with one sheet: a bold header row and the data rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_name(title)
    sheet.append([str(c) for c in columns])
    for row in rows:
        sheet.append([_value(c) for c in row])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    widths = [len(str(c)) for c in columns]
    for row in rows:
        for index, cell in enumerate(row):
            if index >= len(widths):
                widths.append(0)
            longest = max((len(p) for p in str(cell).split("\n")), default=0)
            widths[index] = max(widths[index], longest)
    for index, width in enumerate(widths, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = max(
            MIN_WIDTH, min(width + 2, MAX_WIDTH)
        )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def data_table_to_xlsx(content: dict[str, Any]) -> bytes:
    """Export a Studio data table (``columns`` with ``key``/``label``, ``rows``)."""
    columns = content.get("columns") or []
    keys = [str(c.get("key", "")) for c in columns]
    labels = [str(c.get("label") or c.get("key") or "") for c in columns]
    rows = []
    for row in content.get("rows") or []:
        cells = row.get("cells") or {}
        rows.append([cells.get(key, "") for key in keys])
    return table_to_xlsx(str(content.get("title") or ""), labels, rows)

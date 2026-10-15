"""Office files from structured parts the assistant sends (slides, blocks, sheets).

The parts are already validated and size-capped by the caller; every text still
passes through `clean_text` (no control characters, which Office files reject).
"""

import io
import re
from collections.abc import Sequence
from typing import Any, Final

from app.modules.notebooks.application.export.pptx import slides_to_pptx
from app.modules.notebooks.application.export.xlsx import (
    MAX_WIDTH,
    MIN_WIDTH,
    _sheet_name,
)

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# C0 controls except tab and newline, DEL and C1 controls (XML 1.0 rejects most).
_CONTROL: Final = re.compile(r"[\x00-\x08\x0b-\x1f\x7f-\x9f]")


def clean_text(value: Any) -> str:
    """Return the value as text without control characters."""
    return _CONTROL.sub("", str(value if value is not None else ""))


def _cell(value: Any) -> Any:
    """Numbers stay numbers; text is cleaned (a leading ``=`` is not a formula)."""
    if isinstance(value, int | float) and not isinstance(value, bool):
        return value
    text = clean_text(value)
    return "'" + text if text.startswith("=") else text


def presentation_to_pptx(title: str, slides: Sequence[dict[str, Any]]) -> bytes:
    """A title slide, then one bullet slide per part (notes go to speaker notes)."""
    parts = [
        {
            "layout": "bullets",
            "title": clean_text(s.get("title")),
            "bullets": [clean_text(b) for b in s.get("bullets") or []],
            "notes": clean_text(s.get("notes")),
        }
        for s in slides
    ]
    return slides_to_pptx(clean_text(title), None, parts)


def _docx_table(document: Any, block: dict[str, Any]) -> None:
    """A table with a bold header row."""
    columns = [clean_text(c) for c in block.get("columns") or []]
    rows = [[clean_text(c) for c in r] for r in block.get("rows") or []]
    width = max([len(columns), *(len(r) for r in rows)], default=0)
    if not width:
        return
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for index, row in enumerate([columns, *rows] if columns else rows):
        cells = table.add_row().cells
        for c in range(width):
            cells[c].text = row[c] if c < len(row) else ""
            if index == 0 and columns:
                for run in cells[c].paragraphs[0].runs:
                    run.bold = True


def blocks_to_docx(title: str, blocks: Sequence[dict[str, Any]]) -> bytes:
    """A Word document: the title, then headings, paragraphs, lists and tables."""
    document = Document()
    if clean_text(title).strip():
        document.add_heading(clean_text(title).strip(), level=0)
    for block in blocks:
        kind = block.get("type")
        if kind == "heading":
            level = min(max(int(block.get("level") or 1), 1), 3)
            document.add_heading(clean_text(block.get("text")), level=level)
        elif kind == "paragraph":
            document.add_paragraph(clean_text(block.get("text")))
        elif kind == "bullets":
            for item in block.get("items") or []:
                document.add_paragraph(clean_text(item), style="List Bullet")
        elif kind == "table":
            _docx_table(document, block)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _fit_widths(sheet: Any, rows: list[list[Any]]) -> None:
    """Column widths from the longest value, within MIN_WIDTH..MAX_WIDTH."""
    widths: list[int] = []
    for row in rows:
        for c, value in enumerate(row):
            if c >= len(widths):
                widths.append(0)
            widths[c] = max(widths[c], len(str(value)))
    for c, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(c)].width = max(
            MIN_WIDTH, min(width + 2, MAX_WIDTH)
        )


def sheets_to_xlsx(sheets: Sequence[dict[str, Any]]) -> bytes:
    """A workbook with one sheet per part: bold frozen header row, fitted widths."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    taken: set[str] = set()
    for index, part in enumerate(sheets, start=1):
        name = _sheet_name(clean_text(part.get("name")))
        if name.lower() in taken:
            name = f"{name[:27]} ({index})"
        taken.add(name.lower())
        sheet = workbook.create_sheet(name)
        columns = [clean_text(c) for c in part.get("columns") or []]
        rows = [[_cell(c) for c in r] for r in part.get("rows") or []]
        sheet.append(columns)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        _fit_widths(sheet, [columns, *rows])
    if not workbook.worksheets:
        workbook.create_sheet("Sheet1")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
